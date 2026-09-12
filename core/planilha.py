"""Planilha viva de controle processual.

Gera um .xlsx a partir do acervo, com quatro abas: Acervo, Prazos, Movimentações
e Conferência. A aba Conferência isola justamente o que ainda não foi validado
por humano, para que nenhuma estimativa passe despercebida.

Identidade visual: azul-marinho e dourado do escritório Tiago Neves.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .acervo import Acervo, classificar_urgencia

AZUL_MARINHO = "1B2A4A"
DOURADO = "C6A15B"
BRANCO = "FFFFFF"
CINZA_CLARO = "F2F4F7"

CORES_URGENCIA = {
    "vermelho": "F8D7DA",
    "amarelo": "FFF3CD",
    "verde": "D4EDDA",
}

BORDA = Border(*(Side(style="thin", color="D0D5DD"),) * 4)


def _cabecalho(ws: Worksheet, colunas: list[str], larguras: list[int]) -> None:
    ws.append(colunas)
    for idx, (titulo, largura) in enumerate(zip(colunas, larguras), start=1):
        celula = ws.cell(row=1, column=idx)
        celula.font = Font(name="Calibri", size=11, bold=True, color=BRANCO)
        celula.fill = PatternFill("solid", fgColor=AZUL_MARINHO)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = BORDA
        ws.column_dimensions[get_column_letter(idx)].width = largura
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _estilizar_corpo(ws: Worksheet, primeira_linha: int = 2) -> None:
    for linha in ws.iter_rows(min_row=primeira_linha):
        for celula in linha:
            celula.border = BORDA
            celula.alignment = Alignment(vertical="top", wrap_text=True)


def _aba_acervo(wb: Workbook, acervo: Acervo, hoje: date) -> None:
    ws = wb.active
    ws.title = "Acervo"
    _cabecalho(
        ws,
        ["Nº CNJ", "Cliente", "Polo", "Parte contrária", "Órgão", "Classe",
         "Assunto", "Situação", "Prazos abertos", "Próximo vencimento"],
        [24, 26, 10, 26, 30, 22, 30, 12, 14, 18],
    )
    for proc in sorted(acervo.processos.values(), key=lambda p: p.numero):
        prox = proc.proximo_vencimento()
        ws.append([
            proc.numero, proc.cliente, proc.polo, proc.parte_contraria, proc.orgao,
            proc.classe, proc.assunto, proc.situacao, len(proc.prazos_abertos()),
            prox or "",
        ])
        if prox:
            cor = CORES_URGENCIA[classificar_urgencia(prox, hoje)]
            ws.cell(row=ws.max_row, column=10).fill = PatternFill("solid", fgColor=cor)
    _estilizar_corpo(ws)


def _aba_prazos(wb: Workbook, acervo: Acervo, hoje: date) -> None:
    ws = wb.create_sheet("Prazos")
    _cabecalho(
        ws,
        ["Urgência", "Vencimento estimado", "Dias restantes", "Nº CNJ", "Cliente",
         "Providência", "Evento", "Termo inicial", "Regime", "Dias",
         "Conferido?", "Avisos"],
        [12, 18, 14, 24, 24, 34, 14, 22, 18, 8, 12, 46],
    )
    for proc, prz in acervo.prazos_ate(date(hoje.year + 5, 12, 31)):
        urgencia = classificar_urgencia(prz.vencimento_estimado, hoje)
        faltam = (date.fromisoformat(prz.vencimento_estimado) - hoje).days
        ws.append([
            {"vermelho": "VERMELHO", "amarelo": "AMARELO", "verde": "VERDE"}[urgencia],
            prz.vencimento_estimado, faltam, proc.numero, proc.cliente, prz.descricao,
            prz.evento, prz.termo_inicial, prz.regime, prz.dias,
            "SIM" if prz.conferido_por_humano else "NÃO",
            " | ".join(prz.avisos),
        ])
        fill = PatternFill("solid", fgColor=CORES_URGENCIA[urgencia])
        for col in (1, 2, 3):
            ws.cell(row=ws.max_row, column=col).fill = fill
        if not prz.conferido_por_humano:
            c = ws.cell(row=ws.max_row, column=11)
            c.font = Font(bold=True, color="B42318")
    _estilizar_corpo(ws)


def _aba_movimentacoes(wb: Workbook, acervo: Acervo) -> None:
    ws = wb.create_sheet("Movimentações")
    _cabecalho(ws, ["Data", "Nº CNJ", "Cliente", "Fonte", "Descrição"], [14, 24, 24, 12, 70])
    linhas = [
        (m.data, p.numero, p.cliente, m.fonte, m.descricao)
        for p in acervo.processos.values()
        for m in p.movimentos
    ]
    for linha in sorted(linhas, reverse=True):
        ws.append(list(linha))
    _estilizar_corpo(ws)


def _aba_conferencia(wb: Workbook, acervo: Acervo) -> None:
    ws = wb.create_sheet("Conferência")
    ws["A1"] = "PRAZOS ESTIMADOS PENDENTES DE CONFERÊNCIA HUMANA"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color=AZUL_MARINHO)
    ws["A2"] = (
        "Toda data desta planilha é estimativa produzida por automação. A contagem "
        "definitiva depende de conferência no sistema do tribunal. Enquanto a coluna "
        "\"Conferido?\" estiver como NÃO, o prazo não deve ser tratado como certo."
    )
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 46
    ws["A4"] = "Nº CNJ"
    ws["B4"] = "Cliente"
    ws["C4"] = "Providência"
    ws["D4"] = "Vencimento estimado"
    ws["E4"] = "Base legal aplicada"
    ws["F4"] = "Avisos"
    for col, largura in zip("ABCDEF", [24, 24, 34, 18, 60, 50]):
        c = ws[f"{col}4"]
        c.font = Font(bold=True, color=BRANCO)
        c.fill = PatternFill("solid", fgColor=DOURADO)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[col].width = largura
    linha = 5
    for proc, prz in acervo.nao_conferidos():
        ws.cell(row=linha, column=1, value=proc.numero)
        ws.cell(row=linha, column=2, value=proc.cliente)
        ws.cell(row=linha, column=3, value=prz.descricao)
        ws.cell(row=linha, column=4, value=prz.vencimento_estimado)
        ws.cell(row=linha, column=5, value="\n".join(prz.base_legal))
        ws.cell(row=linha, column=6, value="\n".join(prz.avisos))
        linha += 1
    if linha > 5:
        _estilizar_corpo(ws, primeira_linha=5)
    else:
        ws.cell(row=5, column=1, value="Nenhum prazo pendente de conferência.")


def gerar(
    acervo: Acervo,
    destino: Path | str,
    hoje: date | None = None,
) -> Path:
    """Escreve a planilha de controle processual e devolve o caminho gerado."""
    ref = hoje or date.today()
    wb = Workbook()
    _aba_acervo(wb, acervo, ref)
    _aba_prazos(wb, acervo, ref)
    _aba_movimentacoes(wb, acervo)
    _aba_conferencia(wb, acervo)
    caminho = Path(destino)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho
