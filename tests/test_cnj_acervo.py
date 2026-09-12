"""Testes da numeração CNJ e do acervo."""

from datetime import date

import pytest

from core import cnj
from core.acervo import Acervo, Movimento, Prazo, classificar_urgencia


def numero_valido() -> str:
    d = cnj.gerar_digito("0001234", "2024", "5", "16", "0001")
    return f"0001234-{d}.2024.5.16.0001"


def test_digito_verificador_round_trip():
    num = numero_valido()
    assert cnj.e_valido(num)
    assert cnj.validar(num).formatado == num


def test_digito_incorreto_e_rejeitado():
    num = numero_valido()
    errado = int(num[8:10])
    errado = (errado + 1) % 100
    assert not cnj.e_valido(f"0001234-{errado:02d}.2024.5.16.0001")


def test_aceita_numero_sem_mascara():
    num = numero_valido()
    sem_mascara = num.replace("-", "").replace(".", "")
    assert cnj.validar(sem_mascara).formatado == num


def test_formato_invalido_levanta_erro():
    with pytest.raises(ValueError, match="fora do padrão"):
        cnj.validar("12345")


def test_segmento_e_calendario_sugerido():
    n = cnj.validar(numero_valido())
    assert n.segmento_nome == "Justiça do Trabalho"
    assert n.calendario_sugerido == "justica-do-trabalho"
    assert n.ano == "2024"
    assert n.tribunal == "16"


def test_extrair_ignora_numeros_invalidos():
    num = numero_valido()
    texto = f"Autos {num}; menção inválida 9999999-99.2024.8.10.0001."
    assert cnj.extrair(texto) == [num]


def test_extrair_nao_repete():
    num = numero_valido()
    assert cnj.extrair(f"{num} e de novo {num}") == [num]


def test_movimento_nao_duplica_por_id_externo():
    proc = Acervo().obter_ou_criar(numero_valido())
    m = Movimento(data="2026-09-01", descricao="Intimação", fonte="gmail", id_externo="msg-1")
    assert proc.registrar_movimento(m) is True
    assert proc.registrar_movimento(m) is False
    assert len(proc.movimentos) == 1


def test_prazo_nao_duplica_mesmo_evento():
    proc = Acervo().obter_ou_criar(numero_valido())
    p = Prazo(
        descricao="Contestação", vencimento_estimado="2026-10-01", evento="2026-09-01",
        termo_inicial="dje_disponibilizacao", regime="cpc_dias_uteis", dias=15,
    )
    assert proc.registrar_prazo(p) is True
    assert proc.registrar_prazo(p) is False


def test_status_de_prazo_invalido_e_rejeitado():
    with pytest.raises(ValueError):
        Prazo(
            descricao="x", vencimento_estimado="2026-10-01", evento="2026-09-01",
            termo_inicial="publicacao", regime="cpc_dias_uteis", dias=15,
            status="inventado",
        )


def test_semaforo_de_urgencia():
    hoje = date(2026, 9, 12)
    assert classificar_urgencia("2026-09-14", hoje) == "vermelho"
    assert classificar_urgencia("2026-09-20", hoje) == "amarelo"
    assert classificar_urgencia("2026-10-30", hoje) == "verde"


def test_persistencia_preserva_movimentos_e_prazos(tmp_path):
    acervo = Acervo()
    proc = acervo.obter_ou_criar(numero_valido(), cliente="Cliente Teste")
    proc.registrar_movimento(
        Movimento(data="2026-09-01", descricao="Sentença", fonte="datajud", id_externo="d1")
    )
    proc.registrar_prazo(Prazo(
        descricao="Recurso ordinário", vencimento_estimado="2026-09-25",
        evento="2026-09-01", termo_inicial="dje_disponibilizacao",
        regime="clt_dias_uteis", dias=8,
    ))
    destino = tmp_path / "acervo.json"
    acervo.salvar(destino)

    recarregado = Acervo.carregar(destino)
    proc2 = recarregado.processos[numero_valido()]
    assert proc2.cliente == "Cliente Teste"
    assert proc2.movimentos[0].descricao == "Sentença"
    assert proc2.prazos[0].dias == 8
    assert proc2.proximo_vencimento() == "2026-09-25"


def test_nao_conferidos_lista_apenas_pendentes():
    acervo = Acervo()
    proc = acervo.obter_ou_criar(numero_valido())
    proc.registrar_prazo(Prazo(
        descricao="A", vencimento_estimado="2026-09-25", evento="2026-09-01",
        termo_inicial="publicacao", regime="cpc_dias_uteis", dias=15,
    ))
    proc.registrar_prazo(Prazo(
        descricao="B", vencimento_estimado="2026-09-26", evento="2026-09-02",
        termo_inicial="publicacao", regime="cpc_dias_uteis", dias=15,
        conferido_por_humano=True,
    ))
    assert [p.descricao for _, p in acervo.nao_conferidos()] == ["A"]


def test_planilha_e_gerada_com_as_quatro_abas(tmp_path):
    from openpyxl import load_workbook

    from core import planilha

    acervo = Acervo()
    proc = acervo.obter_ou_criar(numero_valido(), cliente="Cliente Teste", orgao="TRT16")
    proc.registrar_prazo(Prazo(
        descricao="Contestação", vencimento_estimado="2026-09-25", evento="2026-09-01",
        termo_inicial="dje_disponibilizacao", regime="clt_dias_uteis", dias=15,
        avisos=["estimativa"], base_legal=["CLT, art. 775"],
    ))
    destino = planilha.gerar(acervo, tmp_path / "controle.xlsx", hoje=date(2026, 9, 12))
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Acervo", "Prazos", "Movimentações", "Conferência"]
    assert wb["Conferência"]["A5"].value == numero_valido()
